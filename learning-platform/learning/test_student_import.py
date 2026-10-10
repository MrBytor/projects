import csv
import io
import json
from datetime import timedelta
from unittest.mock import patch
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from .models import Course, Enrollment, StudentRegistration, TestAssignment, TestAttempt
from .student_import import HEADERS, MAX_SIZE, apply_import, parse_upload, validate_rows


class StudentImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.teacher = User.objects.create_user('teacher', is_staff=True)
        cls.other_teacher = User.objects.create_user('other-teacher', is_staff=True)
        cls.student = User.objects.create_user('existing', first_name='Original Name', email='old@example.test', password='original-long-password')
        cls.outside = User.objects.create_user('outside', email='private@example.test')
        cls.stats = Course.objects.create(slug='statistics', title='Statistics', teacher=cls.teacher)
        cls.finance = Course.objects.create(slug='corporate-finance', title='Finance', teacher=cls.teacher)
        cls.private = Course.objects.create(slug='private', title='Private course', teacher=cls.other_teacher)
        Enrollment.objects.create(student=cls.student, course=cls.stats)
        Enrollment.objects.create(student=cls.student, course=cls.private)
        StudentRegistration.objects.create(student=cls.student, teacher=cls.teacher, major='Original Major', teaching_group='Original Group')
        cls.work = TestAssignment.objects.create(course=cls.stats, title='Saved work')
        TestAttempt.objects.create(student=cls.student, assignment=cls.work, answers={'preserved': 'draft'})

    def setUp(self):
        self.client.force_login(self.teacher)
        self.url = reverse('import_students')

    def row(self, **values):
        return {'line': 2, **{key: '' for key in HEADERS}, 'username': 'new-student', 'name': 'New Student', **values}

    def csv_file(self, rows, delimiter=',', headers=HEADERS):
        text = io.StringIO()
        writer = csv.writer(text, delimiter=delimiter)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([row.get(key, '') for key in headers])
        return SimpleUploadedFile('students.csv', ('\ufeff' + text.getvalue()).encode('utf8'))

    def xlsx_file(self, rows, headers=HEADERS):
        workbook = Workbook()
        workbook.active.append(headers)
        for row in rows:
            workbook.active.append([row.get(key, '') for key in headers])
        raw = io.BytesIO()
        workbook.save(raw)
        return SimpleUploadedFile('students.xlsx', raw.getvalue())

    def upload(self, rows):
        response = self.client.post(self.url, {'action': 'upload', 'file': self.csv_file(rows)})
        self.assertRedirects(response, self.url)
        return self.client.session['student_import']['token']

    def test_csv_and_excel_support_unicode_and_common_delimiters(self):
        row = self.row(name='Łukasz 王', major='金融', teaching_group='组 1', courses='statistics|corporate-finance')
        for file in [self.csv_file([row]), self.csv_file([row], ';'), self.csv_file([row], '\t'), self.xlsx_file([row])]:
            parsed = parse_upload(file)
            self.assertEqual(parsed[0]['name'], row['name'])
            self.assertEqual(parsed[0]['line'], 2)
            checked = validate_rows(self.teacher, parsed)[0]
            self.assertFalse(checked['errors'])
            self.assertEqual(checked['course_ids'], [self.stats.pk, self.finance.pk])

    def test_parse_errors_limits_headers_formulas_and_invalid_files(self):
        files = [SimpleUploadedFile('students.xls', b'old format'), SimpleUploadedFile('bad.xlsx', b'not a zip'),
                 SimpleUploadedFile('bad.csv', b'\xff\xfe'), SimpleUploadedFile('empty.csv', b''),
                 self.csv_file([self.row()], headers=['name']),
                 self.csv_file([self.row()], headers=['username', 'name', 'name']),
                 self.csv_file([self.row()], headers=['username', 'name', 'is_staff']),
                 self.xlsx_file([self.row(name='=1+1')]),
                 SimpleUploadedFile('extra.csv', b'username,name\na,A,statistics,private\n'),
                 SimpleUploadedFile('large.csv', b'x' * (MAX_SIZE + 1)),
                 self.csv_file([self.row(username=f'row{i}') for i in range(1001)])]
        for file in files:
            with self.subTest(file=file.name), self.assertRaises(ValidationError):
                parse_upload(file)

    def test_preview_changes_nothing_and_apply_creates_hashed_student_login(self):
        before = User.objects.count()
        token = self.upload([self.row(email='new@example.test', major='Finance', teaching_group='F2', courses='statistics')])
        self.assertEqual(User.objects.count(), before)
        preview = self.client.get(self.url)
        self.assertContains(preview, 'Create student')
        self.assertNotContains(preview, 'initial_password')
        self.assertEqual(preview.context['errors'], 0)
        response = self.client.post(self.url, {'action': 'apply', 'token': token})
        self.assertContains(response, 'Import complete')
        self.assertIn('no-store', response['Cache-Control'])
        credentials = response.context['credentials'][0]
        student = User.objects.get(username='new-student')
        self.assertTrue(student.check_password(credentials['password']))
        self.assertNotEqual(student.password, credentials['password'])
        self.assertFalse(student.is_staff)
        self.assertFalse(student.is_superuser)
        self.assertTrue(Enrollment.objects.filter(student=student, course=self.stats).exists())
        self.assertEqual(StudentRegistration.objects.get(student=student, teacher=self.teacher).teaching_group, 'F2')
        self.assertNotIn('student_import', self.client.session)
        self.assertNotIn(credentials['password'], json.dumps(dict(self.client.session)))
        replay = self.client.post(self.url, {'action': 'apply', 'token': token})
        self.assertRedirects(replay, self.url)
        self.assertEqual(User.objects.filter(username='new-student').count(), 1)

    def test_existing_updates_are_additive_and_preserve_work_password_and_blanks(self):
        result = apply_import(self.teacher, [self.row(username='EXISTING', name='Updated Name', courses='corporate-finance')])
        self.assertEqual(result, {'created': 0, 'updated': 1, 'credentials': []})
        self.student.refresh_from_db()
        self.assertEqual(self.student.first_name, 'Updated Name')
        self.assertEqual(self.student.email, 'old@example.test')
        self.assertTrue(self.student.check_password('original-long-password'))
        self.assertEqual(set(Enrollment.objects.filter(student=self.student).values_list('course_id', flat=True)), {self.stats.pk, self.finance.pk, self.private.pk})
        registration = StudentRegistration.objects.get(student=self.student, teacher=self.teacher)
        self.assertEqual(registration.major, 'Original Major')
        self.assertEqual(registration.teaching_group, 'Original Group')
        self.assertEqual(TestAttempt.objects.get(student=self.student).answers, {'preserved': 'draft'})

    def test_duplicate_accounts_emails_roles_and_foreign_courses_are_rejected(self):
        cases = [[self.row(), self.row(username='NEW-STUDENT')],
                 [self.row(email='dup@example.test'), self.row(username='another', email='DUP@example.test')],
                 [self.row(username='outside')], [self.row(username='teacher')],
                 [self.row(email='PRIVATE@example.test')], [self.row(courses='private')],
                 [self.row(username='bad space')], [self.row(email='not-an-email')],
                 [self.row(name='')], [self.row(major='x' * 121)], [self.row(teaching_group='x' * 81)]]
        for rows in cases:
            with self.subTest(rows=rows):
                self.assertTrue(any(row['errors'] for row in validate_rows(self.teacher, rows)))
                with self.assertRaises(ValidationError):
                    apply_import(self.teacher, rows)
        self.assertFalse(User.objects.filter(username='new-student').exists())

    def test_unicode_casefold_and_ambiguous_existing_names_are_not_new_accounts(self):
        User.objects.create_user('Straße', email='Święto@example.test')
        self.assertTrue(validate_rows(self.teacher, [self.row(username='STRASSE')])[0]['errors'])
        self.assertTrue(validate_rows(self.teacher, [self.row(email='święto@example.test')])[0]['errors'])
        User.objects.create_user('Existing')
        self.assertIn('Username is ambiguous', ' '.join(validate_rows(self.teacher, [self.row(username='existing')])[0]['errors']))
        elevated = User.objects.create_user('elevated', is_superuser=True, is_staff=False)
        StudentRegistration.objects.create(student=elevated, teacher=self.teacher)
        self.assertTrue(validate_rows(self.teacher, [self.row(username='elevated')])[0]['errors'])

    def test_all_rows_are_validated_before_saving_and_transactions_roll_back(self):
        original = self.student.first_name
        with self.assertRaises(ValidationError):
            apply_import(self.teacher, [self.row(username='existing', name='Should not save'), self.row(email='bad')])
        self.student.refresh_from_db()
        self.assertEqual(self.student.first_name, original)
        with patch.object(User.objects, 'create_user', side_effect=IntegrityError('concurrent account')):
            with self.assertRaises(IntegrityError):
                apply_import(self.teacher, [self.row(username='existing', name='Should roll back'), self.row()])
        self.student.refresh_from_db()
        self.assertEqual(self.student.first_name, original)

    def test_preview_revalidates_course_access_and_expired_or_bad_tokens(self):
        token = self.upload([self.row(courses='corporate-finance')])
        self.finance.teacher = self.other_teacher
        self.finance.save()
        response = self.client.post(self.url, {'action': 'apply', 'token': token})
        self.assertContains(response, 'No changes were saved')
        self.assertFalse(User.objects.filter(username='new-student').exists())
        self.assertRedirects(self.client.post(self.url, {'action': 'apply', 'token': 'wrong'}), self.url)
        session = self.client.session
        stage = session['student_import']
        stage['created'] = (timezone.now() - timedelta(minutes=31)).timestamp()
        session['student_import'] = stage
        session.save()
        self.assertRedirects(self.client.post(self.url, {'action': 'apply', 'token': token}), self.url)
        self.assertNotIn('student_import', self.client.session)

    def test_preview_pagination_cancel_and_escaped_names(self):
        rows = [self.row(username=f'row{i}', name='<script>bad</script>' if i == 0 else f'Student {i}') for i in range(30)]
        self.upload(rows)
        response = self.client.get(self.url)
        self.assertContains(response, '&lt;script&gt;bad&lt;/script&gt;')
        self.assertEqual(len(response.context['preview']), 25)
        self.assertEqual(len(self.client.get(self.url, {'page': 2}).context['preview']), 5)
        self.assertEqual(response.context['new_count'], 30)
        self.assertRedirects(self.client.post(self.url, {'action': 'cancel'}), self.url)
        self.assertNotIn('student_import', self.client.session)
        self.assertFalse(User.objects.filter(username='row0').exists())

    def test_templates_are_blank_private_and_teacher_only_and_csrf_is_required(self):
        url = reverse('import_template')
        response = self.client.get(url)
        self.assertEqual(response.content.decode('utf-8-sig').strip(), ','.join(HEADERS))
        self.assertEqual(response['Cache-Control'], 'no-store')
        workbook = load_workbook(io.BytesIO(self.client.get(url, {'format': 'xlsx'}).content))
        self.assertEqual(list(workbook.active.values), [tuple(HEADERS)])
        workbook.close()
        csrf = Client(enforce_csrf_checks=True)
        csrf.force_login(self.teacher)
        self.assertEqual(csrf.post(self.url, {'action': 'upload', 'file': self.csv_file([self.row()])}).status_code, 403)
        self.client.force_login(self.student)
        for path in [url, self.url]:
            self.assertEqual(self.client.get(path).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
