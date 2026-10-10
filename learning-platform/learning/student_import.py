"""Bounded uploads, teacher-scoped validation and atomic, additive imports."""
import csv
import io
import re
import secrets
import unicodedata
import zipfile
from collections import Counter
from django import forms
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.db import transaction
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from .models import Course, Enrollment, StudentRegistration
from .students import teacher_students

HEADERS = ['username', 'name', 'email', 'major', 'teaching_group', 'courses']
MAX_ROWS = 1000
MAX_SIZE = 2 * 1024 * 1024


class StudentImportForm(forms.Form):
    file = forms.FileField(label='Student list', widget=forms.FileInput(attrs={'accept': '.csv,.xlsx'}), help_text='CSV (UTF-8) or Excel .xlsx; up to 1,000 students and 2 MB. The first worksheet is used.')


def parse_upload(upload):
    if upload.size > MAX_SIZE:
        raise ValidationError('The file must be 2 MB or smaller.')
    raw = upload.read(MAX_SIZE + 1)
    if len(raw) > MAX_SIZE:
        raise ValidationError('The file must be 2 MB or smaller.')
    try:
        if upload.name.lower().endswith('.csv'):
            text = raw.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            records = []
            for record in csv.reader(io.StringIO(text), dialect):
                records.append(record)
                if len(records) > MAX_ROWS + 1:
                    raise ValidationError('Upload at most 1,000 students per file.')
        elif upload.name.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 2000 or sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValidationError('The workbook is too large when expanded.')
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValidationError('Upload at most 1,000 students per file.')
                if sheet.max_column and sheet.max_column > 20:
                    raise ValidationError('Use the supplied template, with at most 20 columns.')
                records = []
                for record in sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=min(sheet.max_column or 6, 20)):
                    if any(cell.data_type == 'f' for cell in record):
                        raise ValidationError('The workbook contains formulas. Paste values before importing.')
                    values = [str(cell.value) if cell.value is not None else '' for cell in record]
                    records.append(values)
                while records and not any(records[-1]):
                    records.pop()
                if len(records) > MAX_ROWS + 1:
                    raise ValidationError('Upload at most 1,000 students per file.')
            finally:
                workbook.close()
        else:
            raise ValidationError('Choose a .csv or .xlsx file.')
    except (UnicodeDecodeError, csv.Error, zipfile.BadZipFile, OSError, ValueError, KeyError, IndexError, ParseError, DefusedXmlException, InvalidFileException) as error:
        raise ValidationError('The file could not be read. Use the CSV or Excel template.') from error
    if not records:
        raise ValidationError('The file is empty.')
    headers = [v.strip().lower().lstrip('\ufeff') for v in records[0]]
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) != len(set(headers)) or any(h not in HEADERS for h in headers):
        raise ValidationError('Use unique template column names: ' + ', '.join(HEADERS) + '.')
    if not {'username', 'name'}.issubset(headers):
        raise ValidationError('The username and name columns are required.')
    rows = []
    for line, record in enumerate(records[1:], 2):
        if not any(v.strip() for v in record):
            continue
        if len(record) > len(headers) and any(v.strip() for v in record[len(headers):]):
            raise ValidationError(f'Row {line} has extra values. Put course slugs in one cell separated by |.')
        rows.append({'line': line, **{key: record[index].strip() if index < len(record) else '' for index, key in enumerate(headers)}})
    if not rows:
        raise ValidationError('Add at least one student below the header.')
    return rows


def validate_rows(teacher, raw_rows):
    courses = {c.slug: c for c in Course.objects.filter(teacher=teacher)}
    own = set(teacher_students(teacher).values_list('pk', flat=True))
    usernames = [unicodedata.normalize('NFKC', row.get('username', '').strip()) for row in raw_rows]
    duplicates = Counter(u.casefold() for u in usernames)
    emails = [row.get('email', '').strip() for row in raw_rows if row.get('email', '').strip()]
    duplicate_emails = Counter(email.casefold() for email in emails)
    users_by_name, users_by_email = {}, {}
    # Database LOWER() is ASCII-only on SQLite. Compare Unicode account names in
    # Python so a differently cased or normalised upload cannot evade duplicates.
    for user in User.objects.values('id', 'username', 'email', 'is_staff', 'is_superuser').iterator():
        name_key = unicodedata.normalize('NFKC', user['username']).casefold()
        email_key = user['email'].casefold()
        if name_key in duplicates:
            users_by_name.setdefault(name_key, []).append(user)
        if email_key in duplicate_emails:
            users_by_email.setdefault(email_key, []).append(user)
    checked = []
    for raw, username in zip(raw_rows, usernames):
        row = {'line': raw['line'], **{key: raw.get(key, '').strip() for key in HEADERS}, 'errors': [], 'course_ids': [], 'course_titles': []}
        row['username'] = username
        for field, value in [('username', username), ('first_name', row['name']), ('email', row['email'])]:
            try:
                User._meta.get_field(field).clean(value, None)
            except ValidationError as error:
                row['errors'].append(f'{"Name" if field == "first_name" else field.title()}: ' + ' '.join(error.messages))
        if not row['name']:
            row['errors'].append('Name is required.')
        for field, limit in [('major', 120), ('teaching_group', 80)]:
            if len(row[field]) > limit:
                row['errors'].append(f'{field.replace("_", " ").title()} must be {limit} characters or fewer.')
        if duplicates[username.casefold()] > 1:
            row['errors'].append('Duplicate username in this file.')
        users = users_by_name.get(username.casefold(), [])
        existing = users[0] if len(users) == 1 else None
        if len(users) > 1:
            row['errors'].append('Username is ambiguous; use a unique account name.')
        elif existing and (existing['is_staff'] or existing['is_superuser'] or existing['id'] not in own):
            row['errors'].append('This username belongs to an account outside your student directory.')
        row['existing_id'] = existing['id'] if existing else None
        if row['email']:
            if duplicate_emails[row['email'].casefold()] > 1:
                row['errors'].append('Duplicate email in this file.')
            if any(u['id'] != row['existing_id'] for u in users_by_email.get(row['email'].casefold(), [])):
                row['errors'].append('This email is already associated with another account.')
        row['action'] = 'Update existing student' if existing else 'Create student'
        for slug in dict.fromkeys(filter(None, re.split(r'[|;]', row['courses']))):
            slug = slug.strip()
            if slug not in courses:
                row['errors'].append(f'Unknown or unavailable course: {slug}.')
            else:
                row['course_ids'].append(courses[slug].pk)
                row['course_titles'].append(courses[slug].title)
        checked.append(row)
    return checked


@transaction.atomic
def apply_import(teacher, raw_rows):
    rows = validate_rows(teacher, raw_rows)
    if any(row['errors'] for row in rows):
        raise ValidationError('The student list changed or contains errors. Review it again before importing.')
    credentials = []
    for row in rows:
        if row['existing_id']:
            student = User.objects.get(pk=row['existing_id'])
            student.first_name = row['name']
            if row['email']:
                student.email = row['email']
            student.save(update_fields=['first_name', 'email'])
        else:
            password = 'Start-' + secrets.token_urlsafe(18)
            student = User.objects.create_user(row['username'], email=row['email'], password=password, first_name=row['name'])
            credentials.append({'name': row['name'], 'username': row['username'], 'password': password})
        registration, _ = StudentRegistration.objects.get_or_create(teacher=teacher, student=student)
        for field in ('major', 'teaching_group'):
            if row[field]:
                setattr(registration, field, row[field])
        registration.save(update_fields=['major', 'teaching_group'])
        Enrollment.objects.bulk_create([Enrollment(student=student, course_id=pk) for pk in row['course_ids']], ignore_conflicts=True)
    return {'created': len(credentials), 'updated': len(rows) - len(credentials), 'credentials': credentials}
